import tagui as r
import os
import requests
import json
import base64
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
import xlrd
import openpyxl
import sys


class Variable:
    NAME_COLUMN = 'เลขที่เอกสาร'
    SHEET_NAME = 'DATA'

    POWERAUTOMATE_URL = 'https://prod-10.southeastasia.logic.azure.com:443/workflows/a5f947c50cc3447db03b3ea1e59e53c0/triggers/manual/paths/invoke?api-version=2016-06-01&sp=%2Ftriggers%2Fmanual%2Frun&sv=1.0&sig=lc3odltTnbhGvIR05Y_hTEiC7i7o_2UufN7qha8RgKI'

    LIST_CONFIG_CODE = ['ae142f216b10bb2cf0e7348749276442', '4d2bb937608a3437da4ccb868abbb247']
    
    FOLDER_PATH = 'result'
    
    def generate_filename():
        return f'summary_result_{datetime.now().strftime("%Y-%m-%d-%H-%M")}.xlsx'

    FILE_NAME = generate_filename()

    FILE_CONFIG = 'config.json'
    # current_folder = os.path.dirname(os.path.abspath(__file__))
    # FILE_CONFIG = os.path.join(current_folder, FILE_CONFIG)

    side = 'Production'

    def wed():
        try:
            if Variable.side.lower() == 'uat':
                return 'https://uat-oasys-x.ofm.co.th'
            elif Variable.side.lower() == 'dev':
                return 'https://dev-oasys-x.ofm.co.th'
            elif Variable.side.lower() == 'production':
                return 'https://oasys-x.officemate.co.th'
            else:
                output_text.insert(tk.END, f'\n###### Not found your side ######\n')
                output_text.see(tk.END)
                root.update_idletasks()
                return 'www.google.com'
            
        except Exception as e:
            output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at Variable.wed ######\n')
            output_text.see(tk.END)
            root.update_idletasks()

    def check_config_file():
        try:
                
            if os.path.exists(Variable.FILE_CONFIG):
            
                with open(Variable.FILE_CONFIG, 'r', encoding='utf-8') as file:
                    config_code = json.load(file)
                    
                    if 'code' in config_code:
                        if config_code['code'] in Variable.LIST_CONFIG_CODE:
                            return config_code['code']
                        else:
                            messagebox.showwarning("Warning", "Config code is not correct!!!")
                            sys.exit(1)
                    else:
                        messagebox.showwarning("Warning", "'code' key is missing from the config file!!!")
                        sys.exit(1)
            else:
                messagebox.showwarning("Warning", "Not found config file!!!")
                sys.exit(1)
                
        except Exception as e:
            messagebox.showwarning("Warning", f"\n{e}\nError at Variable.check_config_file")
            sys.exit(1)
    
    def click_confirms(r):
        try:
            if Variable.side.lower() == 'production':
                r.click('//*[@id="btnValidateDataConfirmDocument"]')

            elif Variable.side.lower() == 'dev' or Variable.side.lower() == 'uat':
                r.click('//*[@id="poCompleteStatusNone"]')
                r.click('/html/body/div[1]/div[2]/div[4]/div[1]/div/div/div[1]/button')
            else:
                pass
            
        except Exception as e:
            output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at Variable.click_confirms ######\n')
            output_text.see(tk.END)
            root.update_idletasks()
    
    def change_mode():
        try:
            if Variable.side.lower() == 'dev':
                Variable.side = "PRODUCTION"
                mode_button.config(text=Variable.side, bg="red")
                side.config(text=f"{Variable.side.title()} 0.2.1")

            else:
                Variable.side = "DEV"
                mode_button.config(text=Variable.side, bg="blue")
                side.config(text=f"{Variable.side.title()} 0.2.1")
        except Exception as e:
            output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at Variable.change_mode ######\n')
            output_text.see(tk.END)
            root.update_idletasks()
    

def read_xls(file_path, sheet_name):
    try:
        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_name(sheet_name)
        
        data = []
        for row_idx in range(sheet.nrows):
            row = sheet.row_values(row_idx)
            data.append(row)

        return data

    except Exception as e:
        output_text.insert(tk.END, f'\n###### {str(e)} from file: {file_path.split("/")[-1]} ######\n###### Error at read_xls ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
        
        return None
        

def read_xlsx(file_path, sheet_name):
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook[sheet_name]
        
        data = []
        for row in sheet.iter_rows(values_only=True):
            data.append(list(row))

        return data
    
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {str(e)} from file: {file_path.split("/")[-1]} ######\n###### Error at read_xlsx ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
         
        return None


def read_excel(file_path, sheet_name):
    try:
        ext = file_path.split('.')[-1]

        if ext == 'xls':
            return read_xls(file_path, sheet_name)
        elif ext == 'xlsx':
            return read_xlsx(file_path, sheet_name)
        else:
            raise ValueError(f"Unsupported file format: {ext}")
        
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at read_excel ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
        
        
def get_data_from_excel():
    global list_so_number, start_out, end_out
    
    try:
        output_text.config(state=tk.NORMAL)
        start_out = output_text.index(tk.INSERT)
        
        folder_path_data = filedialog.askopenfilenames(filetypes=[("Excel files", "*.xlsx *.xls")])
        
        if folder_path_data:
            
            if not folder_path_data:
                output_text.insert(tk.END, f'\n###### Not found data from path ######\n###### Error at get_data_from_excel ######')
                output_text.see(tk.END)
                root.update_idletasks()
                
                list_so_number = None
            
            else:
                list_so_number = []

                name_file = []
                
                for file in folder_path_data:
                    
                    data = read_excel(file, Variable.SHEET_NAME)
                    
                    if data != None:
                        for col in data[0]:
                            if col == Variable.NAME_COLUMN:
                                index = data[0].index(col)
                                
                                so_number = [str(data[i][index]) for i in range(1,len(data)) if data[i][index] not in ['',None]]
                            
                                name_file.append(file.split('/')[-1])
                            
                            else:
                                pass
                        
                        try:        
                            list_so_number.extend(so_number)
                            
                        except:
                            output_text.insert(tk.END, f'\n###### Not have columns "เลขที่เอกสาร" from file: {file.split("/")[-1]} ######\n###### Error at get_data_from_excel ######')
                            output_text.see(tk.END)
                            root.update_idletasks()
                        
                    else:
                        pass
                    
                list_so_number = list(dict.fromkeys(list_so_number))
                
                if list_so_number:
                    files_text = "\n".join(folder_path_data)
                    file_label.config(text=files_text)
                    
                    file_names = ", ".join(name_file)
                    document_numbers = "\n   ".join(list_so_number)
                    
                    output_text.insert(tk.END, f'\n>>>>\tfile name:\t{file_names}\n\tเลขที่เอกสาร')
                    output_text.see(tk.END)
                    root.update_idletasks()
                    
                    for index, item in enumerate(list_so_number, start=1):
                        output_text.insert(tk.END, f'\n {index}\t{item}')
                        output_text.see(tk.END)
                        root.update_idletasks()
                    
                    run.config(bg="green", fg="white", state=tk.NORMAL)
                
                else:
                    run.config(bg="white", fg="white", state=tk.DISABLED)
                    list_so_number = None
                
            output_text.delete(1.0, start_out)
        else:
            messagebox.showwarning("Warning", "No files selected!")
            
    except Exception as e:
        run.config(bg="white", fg="white", state=tk.DISABLED)
        
        file_label.config(text='Please select data again')
    
        output_text.delete(1.0, start_out)
        
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at get_data_from_excel ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
        
        list_so_number = None

    finally:
        output_text.config(state=tk.DISABLED)
        
        end_out = output_text.index(tk.INSERT)
        

def rpa(data_frame):
    global r, exit_button
    
    try:

        web_url = Variable.wed()

        r.init(visual_automation=False, turbo_mode=True)
        
        r.url(f'{web_url}/PersonalInformation')
        
        while not r.present('/html/body/div/div[2]/div[2]/div') or r.present('//*[@id="login_input_submit"]'):
            r.wait(0.1)

        output_text.insert(tk.END, f'\n>>>>>> Start in... <<<<<<')
        output_text.see(tk.END)
        root.update_idletasks()
        
        for i in range(5):
            r.wait(0.5)
            output_text.insert(tk.END, f"\n {i+1}...")
            output_text.see(tk.END)
            root.update_idletasks()
            
        output_text.insert(tk.END, f'\n-------------------------')
        output_text.see(tk.END)
        root.update_idletasks()

        alert = []

        output_text.insert(tk.END, '\n\tเลขที่เอกสาร')
        output_text.see(tk.END)
        root.update_idletasks()
        
        for i, so_num in enumerate(data_frame, start=1):
            
            r.url(f'{web_url}/SO/Detail/{so_num}')
            
            while not r.present('/html/body/div/div[2]/div[2]/div[1]/div/div/div[1]/span'):
                r.wait(0.1)
                
            # ห้ามลบตัวสำคัญ
            r.wait(0.5)
            
            if  r.present('//*[@id="confirmDocument"]') :
                
                r.click('//*[@id="confirmDocument"]')
                
                while True:
                    if r.read('//*[@id="btnValidateDataConfirmDocument"]') == 'ยืนยันการจัดสินค้า':
                        break
                    else:
                        r.wait(0.1)
                    
                r.wait(0.5)
            
                
                while True:
                    if r.present('//*[@id="btnValidateDataConfirmDocument"]'):
                        
                        Variable.click_confirms(r)
        
                        alert.append([so_num, r.read('/html/body/div/div[2]/div[2]/div[1]/div/div/div[1]/span'), '✅ (กด Confirms เรียบร้อย)'])
                        
                        output_text.insert(tk.END, f'\n{i}\t{so_num}\t✅')
                        output_text.see(tk.END)
                        root.update_idletasks()
                        
                        break
                
                    else:
                        alert.append([so_num, r.read('/html/body/div/div[2]/div[2]/div[1]/div/div/div[1]/span'), '⚠️ (กรุณาตรวจสอบ)'])
                        
                        output_text.insert(tk.END, f'\n{i}\t{so_num}\t⚠️')
                        output_text.see(tk.END)
                        root.update_idletasks()
                        
                        break

            elif not r.present('/html/body/div[3]/div/div/div[1]/h4/span'):

                alert.append([so_num, r.read('/html/body/div/div[2]/div[2]/div[1]/div/div/div[1]/span'), '⚠️ (กรุณาตรวจสอบ)'])
                
                output_text.insert(tk.END, f'\n{i}\t{so_num}\t⚠️')
                output_text.see(tk.END)
                root.update_idletasks()               

            else:
        
                alert.append([so_num, 'Not found data', '❌ (ไม่พบข้อมูล)'])
                
                output_text.insert(tk.END, f'\n{i}\t{so_num}\t❌')
                output_text.see(tk.END)
                root.update_idletasks()

        r.url(f'{web_url}/PersonalInformation')
        r.wait(1)

        r.click('/html/body/div/div[1]/header/div/div[4]/div[2]/li/a')

        r.click('//*[@id="btn_SignOut"]')

        COLUMNS = [['เลขที่เอกสาร', 'Status', 'Comment']]
        
        COLUMNS.extend(alert)

        df = COLUMNS
        
        return df
    
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at rpa ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
        return None

    finally:
        exit_button.config(command=exit_program_add_r)


def save_file_excel(data, FOLDER_PATH, FILE_NAME):
    
    try:
        if not os.path.exists(FOLDER_PATH):
            os.makedirs(FOLDER_PATH)
        FILE_PATH = os.path.join(FOLDER_PATH, FILE_NAME)
        
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "DATA"

        for row in data:
            sheet.append(row)

        workbook.save(FILE_PATH)
        
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at save_file_excel ######\n')
        output_text.see(tk.END)
        root.update_idletasks()


def edit_width_columns_excel(FOLDER_PATH, FILE_NAME):

    try:
        FILE_PATH = os.path.join(FOLDER_PATH, FILE_NAME)
        workbook = openpyxl.load_workbook(FILE_PATH)
        worksheet = workbook['DATA']

        for col in worksheet.columns:
            max_length = 0
            column = col[0].column_letter
            
            for cell in col:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            
            column_width = max_length + 2
            worksheet.column_dimensions[column].width = column_width

        workbook.save(FILE_PATH)

    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at edit_width_columns_excel ######\n')
        output_text.see(tk.END)
        root.update_idletasks()


def send_excel_to_powerautomate(df, FOLDER_PATH, FILE_NAME, code):
    
    try:
        FILE_PATH = os.path.join(FOLDER_PATH, FILE_NAME)
        
        result = dict()
        for i in range(1,len(df)):
            if df[i][2] in result:
                result[df[i][2]] = result[df[i][2]]+1
            else:
                result[df[i][2]] = 1
        
        message = [f'\n\t{key} = {value}' for key, value in result.items()]
        
        output_text.insert(tk.END, ''.join(message))
        output_text.see(tk.END)
        root.update_idletasks()

        URL = Variable.POWERAUTOMATE_URL
        
        with open(FILE_PATH, 'rb') as file:
            file_content = base64.b64encode(file.read()).decode('utf-8')

        payload = {
            'filename': FILE_NAME,
            'filecontent': file_content,
            'message': ', '.join(message),
            'code': code
        }

        headers = {'Content-Type': 'application/json'}
        response = requests.post(URL, data=json.dumps(payload), headers=headers)
        
        if response.status_code == 202:
            pass
        else:
            output_text.insert(tk.END, f"\n\t###### send_excel_to_powerautomate not successfully ######\n")
            output_text.see(tk.END)
            root.update_idletasks()
 
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at send_excel_to_powerautomate ######\n')
        output_text.see(tk.END)
        root.update_idletasks()


def exit_program():
    try:
        root.quit()
        root.destroy()
        
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at exit_program ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
    
    
def exit_program_add_r():
    try:
        root.quit()
        root.destroy()
        r.close()
        
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at exit_program_add_r ######\n')
        output_text.see(tk.END)
        root.update_idletasks()


def update_run_state():
    try:
        root.after(1000, update_run_state)
        
    except Exception as e:
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at update_run_state ######\n')
        output_text.see(tk.END)
        root.update_idletasks()
       
       
def main():
    global start_out, end_out
    try:
        output_text.config(state=tk.NORMAL)

        starttime = datetime.now()
        
        output_text.insert(tk.END, f'\n>>>>\tStart Auto Comfirm...')
        output_text.see(tk.END)
        root.update_idletasks()
        
        data = list_so_number
        
        code = Variable.check_config_file()
        
        FOLDER_PATH = Variable.FOLDER_PATH
        FILE_NAME = Variable.FILE_NAME
        
        if type(data) != type(None):

            df = rpa(data)
            
            if type(df) != type(None):

                save_file_excel(df, FOLDER_PATH, FILE_NAME)

                edit_width_columns_excel(FOLDER_PATH, FILE_NAME)

                send_excel_to_powerautomate(df, FOLDER_PATH, FILE_NAME, code)
                
            else:

                pass
            
            endtime = datetime.now()
            
            output_text.insert(tk.END, f'\n>>>>\tTime work: {endtime - starttime}\n---------------------------------------------------------\n')
            output_text.see(tk.END)
            root.update_idletasks()
            
        else:
            output_text.insert(tk.END, f'\n###### Not found data in excel file. ######')
            output_text.see(tk.END)
            root.update_idletasks()
            
        run.config(state=tk.DISABLED, bg = "white", fg="white")
        messagebox.showinfo("Information", "Finish")
        
    except Exception as e: 
        output_text.insert(tk.END, f'\n###### {e} ######\n###### Error at main ######\n')
        output_text.see(tk.END)
        root.update_idletasks()

    finally:
        
        output_text.config(state=tk.DISABLED)
        select_button.config(command= get_data_from_excel)
        start_out = end_out = tk.END
        

if __name__ == "__main__":
    global root, file_label, output_text, exit_button, run, select_button, mode_button
    check_config = Variable.check_config_file()
    
    # Define fonts
    BUTTON_FONT = ("Arial", 12, "bold")
    LABEL_FONT = ("Helvetica", 12, "italic")
    TEXT_FONT = ("Times New Roman", 12)
    SIDE_FONT = ("Helvetica", 16, "bold")
    CODE_FONT = ("Helvetica", 12, "bold")

    root = tk.Tk()
    root.title("Auto confirms system")
    root.geometry("800x600")

    top_frame = tk.Frame(root)
    top_frame.pack(fill=tk.X)

    select_button = tk.Button(top_frame, text="Select Excel Files", command=get_data_from_excel, bg="blue", fg="white", font=BUTTON_FONT)
    select_button.pack(side=tk.BOTTOM, padx=10, pady=0)

    side = tk.Label(top_frame, text=f"{Variable.side.title()} 0.2.1", font=SIDE_FONT, bg='yellow', fg='black')
    side.pack(side=tk.RIGHT, padx=10, pady=10)
    
    file_label = tk.Label(root, text="No files selected", wraplength=0, font=LABEL_FONT)
    file_label.pack(pady=10)

    bottom_frame = tk.Frame(root)
    bottom_frame.pack(fill=tk.BOTH, expand=True)
    bottom_frame.grid_columnconfigure(0, weight=1)
    bottom_frame.grid_rowconfigure(0, weight=1)

    output_text = tk.Text(bottom_frame, height=10, width=60, state=tk.DISABLED, font=TEXT_FONT)
    output_text.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")


    
    scrollbar = tk.Scrollbar(bottom_frame, command=output_text.yview)
    scrollbar.grid(row=0, column=1, sticky="ns")
    output_text['yscrollcommand'] = scrollbar.set

    run = tk.Button(root, text="run", command=main, font=BUTTON_FONT, state=tk.DISABLED, bg="white", fg="black")
    run.pack(side=tk.LEFT, padx=10, pady=10)

    exit_button = tk.Button(root, text="Exit", command=exit_program, bg="red", fg="white", font=BUTTON_FONT)
    exit_button.pack(side=tk.RIGHT, padx=10, pady=10)
    
    ################## DEVELOPER ONLY ##################
    # mode = tk.Frame(top_frame)
    # mode.pack(side=tk.LEFT, padx=10, pady=5)

    # mode_button = tk.Button(mode, text=Variable.side.upper(), command=Variable.change_mode, bg=f"{'red' if Variable.side.lower() != 'dev' else 'blue'}", fg="white", font=BUTTON_FONT)
    # mode_button.pack(side=tk.LEFT)
    ####################################################
    
    root.after(1000, update_run_state)
        
    root.mainloop()


    